"""
Unit & Integration Tests for FastDoc Core Engine
Menguji filter format, scanner tag OpenXML, excel manager, dan batch zip generator.
"""

import io
import os
import zipfile
import unittest
from datetime import date
from docx import Document

from engine.core.filters import format_date, rupiah, ribuan, terbilang, ceklis
from engine.core.scanner import DocxTagScanner, VariableMeta
from engine.core.renderer import TemplateRenderer
from engine.core.excel_parser import ExcelManager
from engine.core.batch import BatchGenerator


class TestFastDocEngine(unittest.TestCase):

    def test_filters(self):
        """Uji seluruh filter kustom bahasa Indonesia"""
        # 1. Format Tanggal
        self.assertEqual(format_date(date(2026, 6, 20)), "20 Juni 2026")
        self.assertEqual(format_date("2019-09-12"), "12 September 2019")

        # 2. Format Rupiah
        self.assertEqual(rupiah(15000000), "Rp 15.500.000,00" if False else "Rp 15.000.000,00")
        self.assertEqual(rupiah(750000.5), "Rp 750.000,50")

        # 3. Ribuan
        self.assertEqual(ribuan(1500000), "1.500.000")

        # 4. Terbilang
        self.assertEqual(terbilang(1500000), "Satu Juta Lima Ratus Ribu Rupiah")
        self.assertEqual(terbilang(25000), "Dua Puluh Lima Ribu Rupiah")

        # 5. Ceklis
        self.assertEqual(ceklis(True), "✓")
        self.assertEqual(ceklis(False), "")

    def test_docx_template_render_and_scan(self):
        """Buat template uji sederhana, scan variabelnya, dan render in-place"""
        # 1. Buat dokumen template uji di memori
        doc = Document()
        doc.add_heading("SURAT KEPUTUSAN RESMI", level=1)
        doc.add_paragraph("Menimbang dan memutuskan bahwa Saudara {{ nama_pegawai }} dengan NIK {{ nik }}.")
        doc.add_paragraph("Diberikan tunjangan sebesar {{ tunjangan | rupiah }} terhitung sejak {{ tgl_sk | format_date }}.")
        doc.add_paragraph("Catatan Evaluasi: {{ catatan_kinerja }}")

        test_template_path = "/tmp/test_template.docx"
        doc.save(test_template_path)

        # 2. Uji Scanner Tag
        scanner = DocxTagScanner(test_template_path)
        schema = scanner.scan()

        self.assertGreaterEqual(schema.total_variables, 5)
        var_keys = [v.key for v in schema.variables]
        self.assertIn("nama_pegawai", var_keys)
        self.assertIn("nik", var_keys)
        self.assertIn("tunjangan", var_keys)
        self.assertIn("tgl_sk", var_keys)
        self.assertIn("catatan_kinerja", var_keys)

        # Cek tipe inferensi
        tunjangan_var = next(v for v in schema.variables if v.key == "tunjangan")
        self.assertEqual(tunjangan_var.data_type, "number")

        tgl_var = next(v for v in schema.variables if v.key == "tgl_sk")
        self.assertEqual(tgl_var.data_type, "date")

        catatan_var = next(v for v in schema.variables if v.key == "catatan_kinerja")
        self.assertTrue(catatan_var.is_ai_candidate)

        # 3. Uji Rendering In-Place
        renderer = TemplateRenderer(test_template_path)
        context = {
            "nama_pegawai": "Budi Santoso",
            "nik": "3312019901820001",
            "tunjangan": 3500000,
            "tgl_sk": "2026-07-01",
            "catatan_kinerja": "Kinerja sangat memuaskan dan berdedikasi tinggi."
        }
        renderer.render(context)
        out_stream = renderer.save_to_stream()

        # Baca kembali output docx untuk memverifikasi penggantian
        out_doc = Document(out_stream)
        full_text = " ".join([p.text for p in out_doc.paragraphs])

        self.assertIn("Budi Santoso", full_text)
        self.assertIn("3312019901820001", full_text)
        self.assertIn("Rp 3.500.000,00", full_text)
        self.assertIn("1 Juli 2026", full_text)
        self.assertIn("Kinerja sangat memuaskan", full_text)

        # Bersihkan file temp
        if os.path.exists(test_template_path):
            os.remove(test_template_path)

    def test_excel_template_and_batch_zip(self):
        """Uji pembuatan template Excel dan generasi batch ZIP"""
        variables = [
            VariableMeta(key="nama_lengkap", display_name="Nama Lengkap", data_type="string"),
            VariableMeta(key="nisn", display_name="NISN", data_type="string"),
            VariableMeta(key="nilai_uang", display_name="Nilai Uang", data_type="number")
        ]

        # 1. Ekspor template Excel
        excel_stream = ExcelManager.generate_template(variables)
        self.assertIsNotNone(excel_stream)

        # 2. Siapkan data simulasi 3 baris
        records = [
            {"nama_lengkap": "Azzam", "nisn": "1555", "nilai_uang": 100000},
            {"nama_lengkap": "Kenzi", "nisn": "1563", "nilai_uang": 200000},
            {"nama_lengkap": "Qiyas", "nisn": "1574", "nilai_uang": 300000},
        ]

        # 3. Buat template docx sederhana
        doc = Document()
        doc.add_paragraph("Nama: {{ nama_lengkap }} | NISN: {{ nisn }} | Saldo: {{ nilai_uang | rupiah }}")
        tpl_stream = io.BytesIO()
        doc.save(tpl_stream)
        tpl_bytes = tpl_stream.getvalue()

        # 4. Generate batch ZIP
        batch_gen = BatchGenerator(tpl_bytes)
        zip_stream = batch_gen.generate_zip(
            records=records,
            filename_pattern="Surat_{{ nama_lengkap }}.docx"
        )

        # 5. Verifikasi isi file ZIP
        with zipfile.ZipFile(zip_stream, "r") as z:
            namelist = z.namelist()
            self.assertEqual(len(namelist), 3)
            self.assertIn("Surat_Azzam.docx", namelist)
            self.assertIn("Surat_Kenzi.docx", namelist)
            self.assertIn("Surat_Qiyas.docx", namelist)

            # Baca salah satu file di dalam zip
            doc_kenzi = Document(io.BytesIO(z.read("Surat_Kenzi.docx")))
            kenzi_text = " ".join([p.text for p in doc_kenzi.paragraphs])
    def test_run_reconciler_fragmented_runs(self):
        """Uji modul RunReconciler dalam menyatukan tag {{ variabel }} yang terpecah di berbagai <w:t>"""
        from engine.core.run_reconciler import RunReconciler
        fragmented_xml = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
        <w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
            <w:body>
                <w:p>
                    <w:r><w:t>Halo Saudara </w:t></w:r>
                    <w:r><w:t>{{</w:t></w:r>
                    <w:proofErr w:type="spellStart"/>
                    <w:r><w:t>nama_</w:t></w:r>
                    <w:proofErr w:type="spellEnd"/>
                    <w:r><w:t>lengkap}}</w:t></w:r>
                    <w:r><w:t>, selamat datang.</w:t></w:r>
                </w:p>
            </w:body>
        </w:document>""".encode("utf-8")

        reconciled = RunReconciler.reconcile_xml(fragmented_xml).decode("utf-8")
        # Pastikan tag {{nama_lengkap}} utuh di dalam satu tag <w:t>
        self.assertIn("{{nama_lengkap}}", reconciled)
        self.assertNotIn("spellStart", reconciled)

    def test_ai_copilot_domains(self):
        """Uji generator narasi offline AICopilotEngine multi-domain"""
        from engine.core.ai_copilot import AICopilotEngine
        # 1. Domain Kurikulum Merdeka PAUD/SD
        res_paud = AICopilotEngine.generate(
            bullet_notes="aktif menyusun lego balok\nhafal surah pendek dan doa\nperlu bimbingan sabar antre",
            domain="kurikulum_merdeka_paud",
            target_subject="Azzam"
        )
        self.assertEqual(res_paud["engine_mode"], "offline_synthesizer")
        self.assertIn("Azzam", res_paud["narrative"])
        self.assertIn("agama", res_paud["narrative"])

        # 2. Domain Evaluasi Kinerja Pegawai
        res_hr = AICopilotEngine.generate(
            bullet_notes="target proyek selesai tepat waktu\ndisiplin hadir pagi hari",
            domain="kinerja_pegawai",
            target_subject="Budi Santoso"
        )
        self.assertIn("Budi Santoso", res_hr["narrative"])
        self.assertIn("kedisiplinan", res_hr["narrative"])

    def test_fuzzy_excel_matching(self):
        """Uji pencocokan cerdas header Excel ke variabel template (fuzzy matching)"""
        target_vars = ["nama_lengkap", "no_induk", "nisn", "ttl", "kelurahan", "pesan_penguatan"]
        self.assertEqual(ExcelManager.fuzzy_match_header("Nama Murid", target_vars), "nama_lengkap")
        self.assertEqual(ExcelManager.fuzzy_match_header("No. Induk", target_vars), "no_induk")
        self.assertEqual(ExcelManager.fuzzy_match_header("Tempat Tanggal Lahir", target_vars), "ttl")
        self.assertEqual(ExcelManager.fuzzy_match_header("Desa / Kelurahan", target_vars), "kelurahan")
        self.assertEqual(ExcelManager.fuzzy_match_header("Catatan Narasi", target_vars), "pesan_penguatan")

    def test_preview_renderer_html(self):
        """Uji pembuatan pratinjau lembar HTML instan dari docx"""
        from engine.core.preview_renderer import PreviewRenderer
        doc = Document()
        doc.add_heading("LAPORAN HASIL BELAJAR", level=1)
        doc.add_paragraph("Nama Murid: {{ nama }}")
        buf = io.BytesIO()
        doc.save(buf)

        html = PreviewRenderer.render_preview_html(buf, {"nama": "Azzam Alano"})
        self.assertIn("fastdoc-preview-sheet", html)
        self.assertIn("Azzam Alano", html)
        self.assertIn("LAPORAN HASIL BELAJAR", html)


if __name__ == "__main__":
    unittest.main()
