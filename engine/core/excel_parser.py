"""
FastDoc Excel Schema & Parser
Menghasilkan template spreadsheet Excel (.xlsx) kosong berbasis variabel
dan membaca file data Excel untuk generasi batch.
"""

import io
from typing import List, Dict, Any, Union, Optional
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from engine.core.scanner import VariableMeta


class ExcelManager:
    """
    Manajer pembuatan template Excel dan pembacaan data batch.
    """

    @staticmethod
    def generate_template(variables: List[VariableMeta], output_stream_or_path: Union[str, io.BytesIO] = None) -> io.BytesIO:
        """
        Menghasilkan file Excel (.xlsx) dengan header sesuai variabel template.
        """
        wb = Workbook()
        ws = wb.active
        ws.title = "Data Dokumen"

        # Gaya Header: Latar Biru Tua Elegan, Teks Putih Tebal
        header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        thin_border = Border(
            left=Side(style='thin', color='D1D5DB'),
            right=Side(style='thin', color='D1D5DB'),
            top=Side(style='thin', color='D1D5DB'),
            bottom=Side(style='thin', color='D1D5DB')
        )

        # Filter variabel: lewati tabel loop kompleks di header utama
        scalar_vars = [v for v in variables if not v.is_loop]

        # Tulis Header
        for col_idx, var in enumerate(scalar_vars, start=1):
            cell = ws.cell(row=1, column=col_idx, value=var.key)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_alignment
            cell.border = thin_border

        # Tulis Baris Contoh (Sample Data)
        sample_row = []
        for var in scalar_vars:
            if var.data_type == "number":
                sample_row.append(1500000)
            elif var.data_type == "date":
                sample_row.append("2026-06-20")
            elif var.data_type == "text_multiline":
                sample_row.append("Catatan deskripsi perkembangan yang sangat baik...")
            else:
                sample_row.append(f"Contoh {var.display_name}")

        for col_idx, val in enumerate(sample_row, start=1):
            cell = ws.cell(row=2, column=col_idx, value=val)
            cell.border = thin_border
            cell.alignment = Alignment(vertical="center")

        # Atur lebar kolom otomatis
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or "")
                if len(val_str) > max_len:
                    max_len = len(val_str)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 15)

        ws.row_dimensions[1].height = 28
        ws.row_dimensions[2].height = 22

        if isinstance(output_stream_or_path, str):
            wb.save(output_stream_or_path)
            return None
        else:
            stream = io.BytesIO()
            wb.save(stream)
            stream.seek(0)
            return stream

    @staticmethod
    def fuzzy_match_header(header: str, target_variables: List[str]) -> str:
        """
        Mencocokkan nama kolom Excel ke nama variabel template secara cerdas (fuzzy matching).
        Contoh: 'Nama Siswa' -> 'nama_lengkap', 'No. Induk' -> 'no_induk', 'Tgl Lahir' -> 'ttl'.
        """
        clean_h = header.strip().lower().replace(".", "").replace("-", " ")
        snake_h = clean_h.replace(" ", "_")

        # 1. Cek kecocokan langsung
        for var in target_variables:
            if snake_h == var.lower() or clean_h == var.lower():
                return var

        # 2. Kamus sinonim semantik umum dokumen resmi
        synonyms = {
            "nama_lengkap": ["nama", "nama murid", "nama siswa", "nama anak", "nama pegawai", "nama karyawan", "nama lengkap"],
            "nama_panggilan": ["panggilan", "nama panggilan", "alias"],
            "no_induk": ["no induk", "nomor induk", "no induk siswa", "nip", "nik", "nis"],
            "nisn": ["nisn", "nomor nisn"],
            "jenis_kelamin": ["jk", "jenis kelamin", "gender", "sex"],
            "ttl": ["ttl", "tempat tanggal lahir", "tgl lahir", "tanggal lahir", "tempat tgl lahir"],
            "agama": ["agama", "kepercayaan"],
            "nama_ayah": ["ayah", "nama ayah", "orang tua ayah"],
            "nama_ibu": ["ibu", "nama ibu", "orang tua ibu"],
            "kelurahan": ["kelurahan", "desa", "desa kelurahan", "alamat desa"],
            "kecamatan": ["kecamatan", "distrik"],
            "kabupaten": ["kabupaten", "kota", "kabupaten kota", "kab"],
            "provinsi": ["provinsi", "prov"],
            "pesan_penguatan": ["pesan", "catatan", "narasi", "evaluasi", "pesan penguatan", "deskripsi"],
        }

        for var, terms in synonyms.items():
            if var in target_variables:
                if any(term == clean_h or term in clean_h for term in terms):
                    return var

        # 3. Substring match
        for var in target_variables:
            var_clean = var.replace("_", " ")
            if var_clean in clean_h or clean_h in var_clean:
                return var

        return snake_h

    @classmethod
    def parse_data(
        cls,
        file_path_or_stream: Union[str, io.BytesIO],
        target_variables: Optional[List[str]] = None
    ) -> List[Dict[str, Any]]:
        """
        Membaca baris-baris data dari file Excel (.xlsx).
        Jika target_variables disediakan, otomatis memetakan nama kolom Excel ke variabel template.
        """
        wb = load_workbook(file_path_or_stream, data_only=True)
        ws = wb.active

        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return []

        # Baris pertama adalah nama-nama kolom/variabel
        header_row = rows[0]
        raw_headers = [str(h).strip() if h is not None else f"col_{i}" for i, h in enumerate(header_row)]

        mapped_headers = []
        for h in raw_headers:
            if target_variables:
                mapped_headers.append(cls.fuzzy_match_header(h, target_variables))
            else:
                mapped_headers.append(h.lower().replace(" ", "_"))

        data_records = []
        for row in rows[1:]:
            # Abaikan baris kosong total
            if not any(row):
                continue

            record = {}
            for col_idx, val in enumerate(row):
                if col_idx < len(mapped_headers):
                    key = mapped_headers[col_idx]
                    raw_key = raw_headers[col_idx]
                    clean_val = "" if val is None else str(val).strip()
                    record[key] = clean_val
                    # Simpan juga di raw key jika berbeda
                    if raw_key != key:
                        record[raw_key] = clean_val

            data_records.append(record)

        return data_records
